import re
from io import BytesIO
import pdfplumber
from docx import Document

SKILLS = ["python", "sql", "postgresql", "machine learning", "deep learning", "nlp", "flask", "django", "pandas", "numpy", "scikit-learn", "tensorflow", "pytorch", "power bi", "tableau", "docker", "aws", "git", "excel", "java", "javascript", "react", "data analysis"]

def extract_text(filename, content):
    ext = filename.rsplit('.', 1)[-1].lower()
    if ext == 'pdf':
        with pdfplumber.open(BytesIO(content)) as pdf:
            return '\n'.join(page.extract_text() or '' for page in pdf.pages)
    if ext == 'docx':
        doc = Document(BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                paragraphs.extend(cell.text for cell in row.cells)
        return '\n'.join(paragraphs)
    raise ValueError('Only PDF and DOCX are supported')

def extract_skills(text):
    normalized = text.lower()
    return sorted(skill for skill in SKILLS if re.search(r'(?<![\w])' + re.escape(skill) + r'(?![\w])', normalized))

def parse_resume(text):
    email = re.search(r'[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}', text)
    phone = re.search(r'(?<!\d)(?:\+91[-\s]?)?[6-9]\d{9}(?!\d)', re.sub(r'[ ()-]', '', text))
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return {'name_guess': lines[0][:100] if lines else 'Not found', 'email': email.group(0) if email else 'Not found', 'phone': phone.group(0) if phone else 'Not found', 'skills': extract_skills(text)}
